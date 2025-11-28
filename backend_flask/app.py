from flask import Flask, request, jsonify
from flask_cors import CORS
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import os
from openpyxl import Workbook, load_workbook
from datetime import datetime

app = Flask(__name__)
CORS(app)

# ------------------ CONFIGURATION ------------------
SENDER_EMAIL = "kongheyd@rknec.edu"
SENDER_PASSWORD = "wpay zkbb otrs ssar"
RECEIVER_EMAIL = "kongheyd@rknec.edu"

EXCEL_FILE = "asha_ventures_data.xlsx"

NEWSLETTER_SHEET = "Newsletter_Subscribers"
CONTACT_SHEET = "Contact_Form_Entries"

NEWSLETTER_COLUMNS = ["Email", "Subscribed_At", "Source"]
CONTACT_COLUMNS = ["Name", "Email", "Organization", "Message", "Submitted_At"]
# ---------------------------------------------------


# ------------- CREATE EXCEL + SHEETS IF MISSING -------------
def ensure_excel_setup():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()

        # First sheet → Newsletter
        ws1 = wb.active
        ws1.title = NEWSLETTER_SHEET
        ws1.append(NEWSLETTER_COLUMNS)

        # Second sheet → Contact
        ws2 = wb.create_sheet(CONTACT_SHEET)
        ws2.append(CONTACT_COLUMNS)

        wb.save(EXCEL_FILE)
        return

    # If file exists, make sure sheets exist
    wb = load_workbook(EXCEL_FILE)

    if NEWSLETTER_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(NEWSLETTER_SHEET)
        ws.append(NEWSLETTER_COLUMNS)

    if CONTACT_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(CONTACT_SHEET)
        ws.append(CONTACT_COLUMNS)

    wb.save(EXCEL_FILE)


# ------------- NEWSLETTER HELPERS -----------------
def is_duplicate_email(email):
    """Checks newsletter duplicate (case insensitive)."""
    ensure_excel_setup()
    wb = load_workbook(EXCEL_FILE)
    ws = wb[NEWSLETTER_SHEET]

    for row in ws.iter_rows(min_row=2, values_only=True):
        existing_email = row[0]
        if existing_email and existing_email.strip().lower() == email.strip().lower():
            wb.close()
            return True

    wb.close()
    return False


def add_newsletter_email(email, source="unknown"):
    ensure_excel_setup()

    if not is_duplicate_email(email):
        wb = load_workbook(EXCEL_FILE)
        ws = wb[NEWSLETTER_SHEET]

        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        ws.append([email, timestamp, source])
        wb.save(EXCEL_FILE)
        wb.close()
        return True

    return False


# ---------------- CONTACT LOGGING -----------------
def add_contact_entry(name, email, organization, message):
    ensure_excel_setup()

    wb = load_workbook(EXCEL_FILE)
    ws = wb[CONTACT_SHEET]

    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    ws.append([name, email, organization, message, timestamp])
    wb.save(EXCEL_FILE)
    wb.close()


# ---------------- CONTACT ROUTE -------------------
@app.route('/contact', methods=['POST'])
def contact():
    data = request.get_json(silent=True) or {}
    name = (data.get('name') or '').strip()
    email = (data.get('email') or '').strip()
    message = (data.get('message') or '').strip()
    organization = (data.get('organization') or '').strip()

    if not name or not email or not message:
        return jsonify({'error': 'name, email and message are required'}), 400

    # ----- Send Email -----
    subject = "New Contact Message"
    body = (
        f" Hello Asha ventures, {name} with Email id {email} "
        f"from organization {organization or 'Not Provided'} "
        f"has Contacted for {message}."
    )

    msg = MIMEMultipart()
    msg['From'] = SENDER_EMAIL
    msg['To'] = RECEIVER_EMAIL
    msg['Subject'] = subject
    msg.attach(MIMEText(body, 'plain'))

    try:
        with smtplib.SMTP('smtp.gmail.com', 587) as server:
            server.starttls()
            server.login(SENDER_EMAIL, SENDER_PASSWORD)
            server.send_message(msg)
    except Exception as e:
        print("Error sending email:", e)
        return jsonify({'error': 'Failed to send email'}), 500

    # ----- Log Contact -----
    add_contact_entry(name, email, organization, message)

    return jsonify({'message': 'Message delivered and logged successfully'}), 200


# -------------- SUBSCRIBE ROUTE -------------------
@app.route('/subscribe', methods=['POST'])
def subscribe():
    data = request.get_json()
    user_email = data.get('email')
    source = data.get('source', "footer-newsletter")

    if not user_email:
        return jsonify({'error': 'Email is required'}), 400

    # ----- Send Email -----
    subject = "New Newsletter Subscription"
    body = f"Hello Asha Ventures,\n\nThe {user_email} has subscribed to our newsletter."

    msg = MIMEMultipart()
    msg['From'] = SENDER_EMAIL
    msg['To'] = RECEIVER_EMAIL
    msg['Subject'] = subject
    msg.attach(MIMEText(body, 'plain'))

    try:
        with smtplib.SMTP('smtp.gmail.com', 587) as server:
            server.starttls()
            server.login(SENDER_EMAIL, SENDER_PASSWORD)
            server.send_message(msg)
    except Exception as e:
        print("Error sending email:", e)
        return jsonify({'error': 'Failed to send email'}), 500

    # ----- Log Newsletter Subscription -----
    added = add_newsletter_email(user_email, source)

    if not added:
        return jsonify({'message': 'Email already exists in the subscriber list'}), 200

    return jsonify({'message': 'Email sent and added to Excel successfully'}), 200


if __name__ == '__main__':
    ensure_excel_setup()
    app.run(host='0.0.0.0', port=5000, debug=True)
