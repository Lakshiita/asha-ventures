# backend_flask/app.py
from flask import Flask, request, jsonify
from flask_cors import CORS
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import os
from openpyxl import Workbook, load_workbook
import logging

app = Flask(__name__)
CORS(app)

# Configure logging for both console and file
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.StreamHandler(),  # Console output
        logging.FileHandler('app.log')  # File output
    ]
)

# ------------------ CONFIGURATION ------------------
SENDER_EMAIL = "kongheyd@rknec.edu"      # Replace with your Gmail
SENDER_PASSWORD = "wpay zkbb otrs ssar"      # Replace with your Gmail App Password
RECEIVER_EMAIL = "kongheyd@rknec.edu"      # Admin notification email
EXCEL_FILE = "subscribed_emails.xlsx"      # Excel file path
SHEET_NAME = "Sheet1"                      # Sheet where emails are stored
COLUMN_NAME = "Subscribed_Emails"          # Column header
# ---------------------------------------------------

def ensure_excel_exists():
    """Create the Excel file if it doesn’t exist, with header."""
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append([COLUMN_NAME])  # Add header row
        wb.save(EXCEL_FILE)

def is_duplicate_email(email):
    """Check if the email already exists in the Excel file."""
    wb = load_workbook(EXCEL_FILE)
    ws = wb[SHEET_NAME]

    for row in ws.iter_rows(min_row=2, values_only=True):  # skip header
        if row[0] and row[0].strip().lower() == email.strip().lower():
            wb.close()
            return True
    wb.close()
    return False

def add_email_to_excel(email):
    """Add email to Excel if not duplicate."""
    ensure_excel_exists()
    if not is_duplicate_email(email):
        wb = load_workbook(EXCEL_FILE)
        ws = wb[SHEET_NAME]
        ws.append([email])
        wb.save(EXCEL_FILE)
        wb.close()
        return True
    return False


@app.route('/subscribe', methods=['POST'])
def subscribe():
    data = request.get_json()
    user_email = data.get('email')
    logging.info(f"Subscription request: {user_email}")

    if not user_email:
        return jsonify({'error': 'Email is required'}), 400

    ensure_excel_exists()

    # ---------------- SEND EMAIL ----------------
    subject = "New Newsletter Subscription"
    body = f"Hello Asha Ventures,\n\nThe {user_email} has subscribed to our newsletter."

    msg = MIMEMultipart()
    msg['From'] = SENDER_EMAIL
    msg['To'] = RECEIVER_EMAIL
    msg['Subject'] = subject
    msg.attach(MIMEText(body, 'plain'))

    try:
        with smtplib.SMTP('smtp.gmail.com', 587) as server:
            server.starttls()
            server.login(SENDER_EMAIL, SENDER_PASSWORD)
            server.send_message(msg)
        logging.info(f"Email sent: {user_email}")
    except Exception as e:
        logging.error(f"Email failed: {e}")
        return jsonify({'error': 'Failed to send email'}), 500

    # ---------------- SAVE TO EXCEL ----------------
    added = add_email_to_excel(user_email)
    logging.info(f"Excel operation complete: {user_email}, added: {added}")

    if not added:
        return jsonify({'message': 'Email already exists in the subscriber list'}), 200

    return jsonify({'message': 'Email sent and added to Excel successfully'}), 200


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
